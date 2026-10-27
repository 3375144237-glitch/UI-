#!/usr/bin/env python3
"""Write a review note and optional reference image to a new Excel workbook."""

from __future__ import annotations

import argparse
from copy import copy
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.drawing.image import Image as SheetImage
from PIL import Image


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--workbook", type=Path, required=True, help="Source .xlsx workbook")
    parser.add_argument("--sheet", required=True, help="Worksheet name")
    parser.add_argument("--cell", help="Cell that receives the note text")
    note_group = parser.add_mutually_exclusive_group()
    note_group.add_argument("--note", help="Note text")
    note_group.add_argument("--note-file", type=Path, help="UTF-8 text file containing the note")
    parser.add_argument("--image", type=Path, help="Optional reference image (PNG or JPEG)")
    parser.add_argument("--image-cell", help="Top-left cell for the reference image; defaults to --cell")
    parser.add_argument("--max-width-px", type=int, default=280, help="Maximum embedded image width (default: 280)")
    parser.add_argument("--max-height-px", type=int, help="Optional maximum embedded image height")
    parser.add_argument("--output", type=Path, required=True, help="Output .xlsx path; source is never overwritten")
    parser.add_argument("--overwrite", action="store_true", help="Allow replacing an existing output file")
    return parser


def main() -> None:
    parser = build_parser()
    args = parser.parse_args()

    source = args.workbook.expanduser().resolve()
    output = args.output.expanduser().resolve()
    if not source.is_file():
        parser.error(f"Workbook not found: {source}")
    if source == output:
        parser.error("Output must differ from the source workbook")
    if source.suffix.lower() != ".xlsx" or output.suffix.lower() != ".xlsx":
        parser.error("Only .xlsx workbooks are supported")
    if output.exists() and not args.overwrite:
        parser.error(f"Output already exists; pass --overwrite to replace it: {output}")

    note = args.note
    if args.note_file:
        note_path = args.note_file.expanduser().resolve()
        if not note_path.is_file():
            parser.error(f"Note file not found: {note_path}")
        note = note_path.read_text(encoding="utf-8")

    if note is None and args.image is None:
        parser.error("Provide --note/--note-file, --image, or both")
    if note is not None and not args.cell:
        parser.error("--cell is required when note text is provided")
    if args.image and args.max_width_px <= 0:
        parser.error("--max-width-px must be greater than zero")
    if args.max_height_px is not None and args.max_height_px <= 0:
        parser.error("--max-height-px must be greater than zero")

    image_path: Path | None = None
    image_cell = args.image_cell or args.cell
    if args.image:
        image_path = args.image.expanduser().resolve()
        if not image_path.is_file():
            parser.error(f"Image not found: {image_path}")
        if image_path.suffix.lower() not in {".png", ".jpg", ".jpeg"}:
            parser.error("Reference image must be PNG or JPEG")
        if not image_cell:
            parser.error("--image-cell is required when --cell is not provided")

    try:
        workbook = load_workbook(source)
        if args.sheet not in workbook.sheetnames:
            parser.error(f"Worksheet not found: {args.sheet}")
        sheet = workbook[args.sheet]

        if note is not None:
            cell = sheet[args.cell]
            cell.value = note
            alignment = copy(cell.alignment)
            alignment.wrap_text = True
            alignment.vertical = "top"
            cell.alignment = alignment

        if image_path:
            with Image.open(image_path) as reference:
                image_width, image_height = reference.size
                if reference.format not in {"PNG", "JPEG"}:
                    parser.error("Reference image must contain PNG or JPEG data")
            scale = min(1.0, args.max_width_px / image_width)
            if args.max_height_px:
                scale = min(scale, args.max_height_px / image_height)
            embedded = SheetImage(str(image_path))
            embedded.width = max(1, round(image_width * scale))
            embedded.height = max(1, round(image_height * scale))
            embedded.anchor = image_cell
            sheet.add_image(embedded)

        output.parent.mkdir(parents=True, exist_ok=True)
        workbook.save(output)
    except (OSError, ValueError, KeyError) as error:
        parser.error(str(error))

    print(f"Saved review workbook: {output}")


if __name__ == "__main__":
    main()
