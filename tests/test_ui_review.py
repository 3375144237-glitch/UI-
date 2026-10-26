import json
import sys
import tempfile
import argparse
import unittest
from pathlib import Path

from PIL import Image, ImageDraw
from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))
import ui_review


class NormalizeTests(unittest.TestCase):
    def test_rem_resizes_to_720_and_keeps_aspect_ratio(self):
        image = Image.new("RGB", (1440, 2880), "white")
        result, info = ui_review.normalize_device(image, "rem", None, None, False)
        self.assertEqual(result.size, (720, 1440))
        self.assertEqual(info["mode"], "REM")

    def test_px_already_two_x_is_not_resized(self):
        image = Image.new("RGB", (780, 1600), "white")
        result, info = ui_review.normalize_device(image, "px", 780, None, True)
        self.assertEqual(result.size, (780, 1600))
        self.assertEqual(info["scale"], 1.0)

    def test_px_scale_must_match_confirmed_target_width(self):
        image = Image.new("RGB", (1080, 2400), "white")
        result, _ = ui_review.normalize_device(image, "px", 720, 3.0, False)
        self.assertEqual(result.size, (720, 1600))
        with self.assertRaises(ValueError):
            ui_review.normalize_device(image, "px", 780, 3.0, False)

    def test_already_two_x_width_mismatch_is_rejected(self):
        image = Image.new("RGB", (720, 1200), "white")
        with self.assertRaises(ValueError):
            ui_review.normalize_device(image, "px", 780, None, True)


class ComparisonTests(unittest.TestCase):
    def test_matching_images_produce_candidate_box(self):
        design = Image.new("RGB", (720, 360), "white")
        device = design.copy()
        ImageDraw.Draw(device).rectangle((100, 100, 125, 125), fill="black")
        _, heat, candidates = ui_review.difference_candidates(design, device, threshold=20, min_area=20)
        self.assertEqual(heat.size, design.size)
        self.assertEqual(len(candidates), 1)
        self.assertGreaterEqual(candidates[0]["bbox"]["width"], 20)

    def test_different_dimensions_do_not_create_pixel_candidates(self):
        with tempfile.TemporaryDirectory() as temp:
            session = Path(temp)
            design = Image.new("RGB", (720, 100), "white")
            device = Image.new("RGB", (780, 100), "white")
            comparable, candidates = ui_review._save_preview(session, design, device)
            self.assertFalse(comparable)
            self.assertEqual(candidates, [])
            self.assertTrue((session / "side-by-side.png").exists())
            self.assertTrue((session / "overlay-left.png").exists())
            with Image.open(session / "overlay.png") as overlay:
                self.assertEqual(overlay.size, design.size)

    def test_red_layer_is_red_tinted_at_half_opacity(self):
        layer = ui_review._red_layer(Image.new("RGB", (2, 2), "white"))
        self.assertEqual(layer.mode, "RGBA")
        self.assertEqual(layer.getpixel((0, 0)), (255, 0, 0, 128))

    def test_overlay_alignments_and_vertical_offset(self):
        with tempfile.TemporaryDirectory() as temp:
            session = Path(temp)
            design = Image.new("RGB", (100, 40), "white")
            device = Image.new("RGB", (20, 10), "white")
            metadata = ui_review.save_overlay_previews(session, design, device, "center", 5)

            self.assertEqual(metadata["x_offsets"], {"left": 0, "right": 80, "center": 40})
            with Image.open(session / "overlay-left.png") as left:
                self.assertEqual(left.getpixel((0, 5)), (255, 127, 127))
                self.assertEqual(left.getpixel((0, 4)), (255, 255, 255))
                self.assertEqual(left.getpixel((20, 5)), (255, 255, 255))
            with Image.open(session / "overlay-right.png") as right:
                self.assertNotEqual(right.getpixel((99, 5)), (255, 255, 255))
                self.assertEqual(right.getpixel((79, 5)), (255, 255, 255))
            with Image.open(session / "overlay-center.png") as center:
                self.assertNotEqual(center.getpixel((40, 5)), (255, 255, 255))
                self.assertEqual(center.getpixel((39, 5)), (255, 255, 255))

    def test_overlay_does_not_resize_device_when_widths_differ(self):
        with tempfile.TemporaryDirectory() as temp:
            session = Path(temp)
            design = Image.new("RGB", (100, 40), "white")
            device = Image.new("RGB", (120, 20), "white")
            metadata = ui_review.save_overlay_previews(session, design, device, "right", -3)
            self.assertEqual(metadata["x_offsets"]["right"], -20)
            with Image.open(session / "device-red-layer.png") as layer:
                self.assertEqual(layer.size, device.size)
            with Image.open(session / "overlay.png") as overlay:
                self.assertEqual(overlay.size, design.size)

    def test_crop_clamps_to_image_bounds_and_annotates(self):
        with tempfile.TemporaryDirectory() as temp:
            destination = Path(temp) / "crop.jpg"
            image = Image.new("RGB", (100, 100), "white")
            ui_review.crop_annotated(image, {"x": -4, "y": 2, "width": 30, "height": 20}, "R-001", destination)
            self.assertTrue(destination.exists())
            with Image.open(destination) as result:
                self.assertLessEqual(result.width, 54)

    def test_chinese_issue_number_renders_with_available_font(self):
        font = ui_review.annotation_font(14)
        bounds = font.getbbox("差异-001")
        self.assertGreater(bounds[2], 0)


class WorkbookTests(unittest.TestCase):
    def test_report_embeds_crop_and_preserves_manual_fields(self):
        with tempfile.TemporaryDirectory() as temp:
            session = Path(temp)
            Image.new("RGB", (720, 240), "white").save(session / "device-prepared.png")
            review = {
                "schema_version": 1,
                "page": "设置页/默认态",
                "device": "Test device",
                "mode": "REM",
                "device_image": {"prepared": "device-prepared.png", "size": [720, 240]},
                "findings": [{
                    "id": "R-001",
                    "description": "按钮偏下",
                    "priority": "P1",
                    "status": "待确认",
                    "remarks": "",
                    "bbox": {"x": 20, "y": 30, "width": 40, "height": 20},
                    "category": "布局与对齐",
                    "affected_pages": ["首页/默认态", "首页/长文案"],
                }],
            }
            (session / "review.json").write_text(json.dumps(review, ensure_ascii=False), encoding="utf-8")
            report = ui_review.generate_report(session, review)
            workbook = load_workbook(report)
            sheet = workbook["验收记录"]
            self.assertEqual([sheet.cell(1, col).value for col in range(1, 7)], ui_review.VISIBLE_HEADERS)
            self.assertEqual(len(sheet._images), 1)
            self.assertTrue(sheet.column_dimensions["G"].hidden)
            info_sheet = workbook["验收信息"]
            self.assertEqual(info_sheet.cell(5, 2).value, "处理后设备截图 @2x 像素坐标")
            detail_sheet = workbook["问题详情"]
            self.assertEqual(detail_sheet.cell(1, 2).value, "问题类别")
            self.assertEqual(detail_sheet.cell(2, 2).value, "布局与对齐")
            self.assertEqual(detail_sheet.cell(2, 3).value, "首页/默认态、首页/长文案")
            self.assertEqual(detail_sheet.cell(1, 10).value, "证据图")
            sheet.cell(2, 3, "手动修订后的描述")
            sheet.cell(2, 4, "P2")
            sheet.cell(2, 5, "修复中")
            sheet.cell(2, 6, "保留备注")
            workbook.save(report)

            ui_review.generate_report(session, review)
            updated = load_workbook(report)["验收记录"]
            self.assertEqual(updated.cell(2, 3).value, "手动修订后的描述")
            self.assertEqual(updated.cell(2, 4).value, "P2")
            self.assertEqual(updated.cell(2, 5).value, "修复中")
            self.assertEqual(updated.cell(2, 6).value, "保留备注")

    def test_previous_report_fields_carry_to_a_new_review_folder(self):
        with tempfile.TemporaryDirectory() as temp:
            base = Path(temp)
            previous = base / "previous"
            current = base / "current"
            previous.mkdir()
            current.mkdir()
            Image.new("RGB", (720, 240), "white").save(previous / "device-prepared.png")
            Image.new("RGB", (720, 240), "white").save(current / "device-prepared.png")
            finding = {
                "id": "stable-button-spacing",
                "description": "按钮偏下",
                "priority": "P1",
                "status": "待修复",
                "remarks": "研发已确认",
                "bbox": {"x": 20, "y": 30, "width": 40, "height": 20},
            }
            review = {
                "schema_version": 1,
                "page": "首页",
                "device": "Test device",
                "mode": "PX",
                "device_image": {"prepared": "device-prepared.png", "size": [720, 240]},
                "findings": [finding],
            }
            old_report = ui_review.generate_report(previous, review)
            old_book = load_workbook(old_report)
            old_sheet = old_book["验收记录"]
            old_sheet.cell(2, 5, "修复中")
            old_sheet.cell(2, 6, "下轮继续跟进")
            old_book.save(old_report)

            new_report = ui_review.generate_report(current, review, str(old_report))
            new_sheet = load_workbook(new_report)["验收记录"]
            self.assertEqual(new_sheet.cell(2, 5).value, "修复中")
            self.assertEqual(new_sheet.cell(2, 6).value, "下轮继续跟进")

    def test_batch_report_keeps_duplicate_issue_ids_separate_by_screenshot(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp) / "batch"
            sessions = [root / "01-page-a", root / "02-page-b", root / "03-clean"]
            for index, session in enumerate(sessions):
                session.mkdir(parents=True)
                Image.new("RGB", (720, 240), "white").save(session / "device-prepared.png")
                Image.new("RGB", (720, 240), "white").save(session / "side-by-side.png")
                Image.new("RGB", (720, 240), "white").save(session / "overlay.png")
                review = {
                    "schema_version": 1,
                    "page": f"页面{index + 1}",
                    "device": "iOS",
                    "design": {"source_name": f"design-{index + 1}.png"},
                    "device_image": {
                        "prepared": "device-prepared.png",
                        "original": "inputs/device-original.png",
                    },
                    "findings": [] if index == 2 else [{
                            "id": "NAV-01",
                            "description": "底部导航间距不一致",
                            "category": "间距",
                            "priority": "P0",
                            "status": "待确认",
                            "bbox": {"x": 20 + index * 10, "y": 30, "width": 40, "height": 20},
                        }],
                }
                if index == 2:
                    review["review_status"] = "已验收，未发现明确差异"
                (session / "review.json").write_text(json.dumps(review, ensure_ascii=False), encoding="utf-8")
            (root / "配对与进度.md").write_text(
                "| 序号 | 设备截图 | Figma 配对 | 状态 | 结果/说明 |\n"
                "|---|---|---|---|---|\n"
                "| 01 | `IMG_1.png` | 1:1 | 已验收 | 01-page-a/ |\n"
                "| 02 | `IMG_2.png` | 1:2 | 已验收 | 02-page-b/ |\n"
                "| 03 | `IMG_3.png` | 1:3 | 已验收，未发现明确差异 | 03-clean/ |\n"
                "| 04 | `IMG_4.png` | 1:4 | 已配对 | 04-page/，问题见 P14-01 |\n",
                encoding="utf-8",
            )

            report = ui_review.generate_batch_report(root)
            workbook = load_workbook(report)
            sheet = workbook["验收记录"]
            self.assertEqual([sheet.cell(1, col).value for col in range(1, 7)], ui_review.VISIBLE_HEADERS)
            self.assertEqual(sheet.max_row, 3)
            self.assertEqual(len(sheet._images), 2)
            self.assertEqual(sheet.cell(2, 3).value, "[NAV-01] 底部导航间距不一致 · 来源截图：01-page-a/IMG_1.png")
            self.assertEqual(sheet.cell(3, 3).value, "[NAV-01] 底部导航间距不一致 · 来源截图：02-page-b/IMG_2.png")
            self.assertNotEqual(sheet.cell(2, 7).value, sheet.cell(3, 7).value)
            details = workbook["问题详情"]
            self.assertEqual(details.max_row, 3)
            self.assertEqual(details.cell(2, 5).value, "01-page-a/IMG_1.png")
            self.assertEqual(details.cell(3, 5).value, "02-page-b/IMG_2.png")
            coverage = workbook["截图清单"]
            self.assertEqual(coverage.max_row, 5)
            self.assertEqual(coverage.cell(2, 3).value, "01-page-a/IMG_1.png")
            self.assertEqual(coverage.cell(4, 7).value, "已验收，未发现明确差异")
            self.assertTrue(coverage.cell(4, 8).hyperlink)
            self.assertTrue(coverage.cell(4, 9).hyperlink)
            self.assertEqual(coverage.cell(4, 9).value, "03-clean/overlay.png")
            self.assertEqual(coverage.cell(5, 3).value, "IMG_4.png")
            self.assertEqual(coverage.cell(5, 6).value, 1)
            self.assertIn("P14-01", coverage.cell(5, 7).value)
            self.assertIn("没有原始设备图", coverage.cell(5, 11).value)

            sheet.cell(2, 5, "修复中")
            sheet.cell(2, 6, "保留备注")
            workbook.save(report)
            ui_review.generate_batch_report(root)
            updated = load_workbook(report)["验收记录"]
            self.assertEqual(updated.cell(2, 5).value, "修复中")
            self.assertEqual(updated.cell(2, 6).value, "保留备注")


class WorkflowTests(unittest.TestCase):
    def test_compare_command_preserves_sources_and_generates_report(self):
        with tempfile.TemporaryDirectory() as temp:
            base = Path(temp)
            design_path = base / "design.png"
            device_path = base / "device.png"
            output = base / "review"
            design = Image.new("RGB", (720, 360), "white")
            device = design.copy()
            ImageDraw.Draw(device).rectangle((250, 180, 275, 205), fill="black")
            design.save(design_path)
            device.save(device_path)
            original_device_bytes = device_path.read_bytes()
            args = argparse.Namespace(
                design=str(design_path),
                device=str(device_path),
                out=str(output),
                mode="rem",
                target_width=None,
                screenshot_scale=None,
                already_2x=False,
                device_name="Synthetic phone",
                page="Home/default",
                previous_report=None,
                design_y=None,
            )
            session = ui_review.prepare_session(args)
            review = json.loads((session / "review.json").read_text(encoding="utf-8"))
            self.assertTrue(review["automatic_pixel_comparison"])
            self.assertGreaterEqual(len(review["findings"]), 1)
            self.assertTrue((session / "report.xlsx").exists())
            self.assertTrue((session / "device-red-layer.png").exists())
            self.assertTrue((session / "overlay-left.png").exists())
            self.assertTrue((session / "overlay-right.png").exists())
            self.assertTrue((session / "overlay-center.png").exists())
            self.assertEqual(device_path.read_bytes(), original_device_bytes)
            self.assertTrue((session / "inputs" / "device-original.png").exists())

            self.assertEqual(
                ui_review.main(["align", "--session", str(session), "--alignment", "right", "--offset-y", "18"]),
                0,
            )
            updated = json.loads((session / "review.json").read_text(encoding="utf-8"))
            self.assertEqual(updated["overlay"]["alignment"], "right")
            self.assertEqual(updated["overlay"]["offset_y"], 18)
            self.assertEqual(updated["findings"], review["findings"])
            info = load_workbook(session / "report.xlsx")["验收信息"]
            self.assertEqual(info.cell(9, 2).value, "right")
            self.assertEqual(info.cell(10, 2).value, 18)

    def test_long_design_is_cropped_at_matched_y_to_device_viewport_height(self):
        with tempfile.TemporaryDirectory() as temp:
            base = Path(temp)
            design_path = base / "long-design.png"
            device_path = base / "middle-screen.png"
            output = base / "review"
            design = Image.new("RGB", (720, 900), "white")
            device = Image.new("RGB", (720, 300), "white")
            ImageDraw.Draw(design).rectangle((80, 340, 110, 370), fill="black")
            ImageDraw.Draw(device).rectangle((80, 90, 110, 120), fill="black")
            design.save(design_path)
            device.save(device_path)
            args = argparse.Namespace(
                design=str(design_path),
                device=str(device_path),
                out=str(output),
                mode="rem",
                target_width=None,
                screenshot_scale=None,
                already_2x=False,
                device_name="Synthetic phone",
                page="Long page/middle",
                previous_report=None,
                design_y=250,
            )

            session = ui_review.prepare_session(args)
            review = json.loads((session / "review.json").read_text(encoding="utf-8"))
            self.assertEqual(review["design"]["source_size"], [720, 900])
            self.assertEqual(review["design"]["reference_region"], {"x": 0, "y": 250, "width": 720, "height": 300})
            self.assertTrue(review["automatic_pixel_comparison"])
            with Image.open(session / "design-prepared.png") as prepared_design:
                self.assertEqual(prepared_design.size, (720, 300))
            self.assertEqual(len(review["findings"]), 0)
            info = load_workbook(session / "report.xlsx")["验收信息"]
            self.assertEqual(info.cell(6, 1).value, "设计稿参考区段（@2x坐标）")
            self.assertIn("'y': 250", info.cell(6, 2).value)

    def test_long_design_crop_out_of_bounds_is_rejected(self):
        with tempfile.TemporaryDirectory() as temp:
            base = Path(temp)
            design_path = base / "long-design.png"
            device_path = base / "screen.png"
            Image.new("RGB", (720, 500), "white").save(design_path)
            Image.new("RGB", (720, 300), "white").save(device_path)
            args = argparse.Namespace(
                design=str(design_path),
                device=str(device_path),
                out=str(base / "review"),
                mode="rem",
                target_width=None,
                screenshot_scale=None,
                already_2x=False,
                device_name="Synthetic phone",
                page="Long page/end",
                previous_report=None,
                design_y=250,
            )
            with self.assertRaisesRegex(ValueError, "超出设计稿高度"):
                ui_review.prepare_session(args)


if __name__ == "__main__":
    unittest.main()
